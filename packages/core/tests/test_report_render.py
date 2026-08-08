"""Dosya raporu oluşturucularının testleri — PDF + Word (E1.5).

En kritik iddia **Türkçe karakterlerin geri okunabilmesidir**: reportlab'ın varsayılan
Type1 fontları ``ş ğ ı İ`` taşımaz ve rapor sessizce kutucuklarla dolardı. Test, üretilen
PDF'ten metni yeniden çıkarıp harflerin yerinde olduğunu doğrular (bkz. ADR-0008).

İkinci iddia: **elenen eşleşme rapordan düşmez** ve gerekçesiyle birlikte basılır.
Üçüncüsü: **TCKN maskeli** çıkar — dosya kurumdan dışarı giden bir belgedir.
"""

from __future__ import annotations

import io
import uuid
from datetime import UTC, datetime

import pytest
from docx import Document as DocxDocument
from pypdf import PdfReader

from aegis_core.reports.dossier import (
    DiscrepancyRow,
    DocumentRow,
    DossierData,
    MatchRow,
    ObservationRow,
    ProfileField,
    RelatedParty,
    ScreeningSection,
    dataset_version_rows,
)
from aegis_core.reports.render import (
    render_dossier_docx,
    render_dossier_html,
    render_dossier_pdf,
)
from aegis_core.reports.risk import (
    DiscrepancyView,
    MatchView,
    ScreeningView,
    assess_risk,
)

NOW = datetime(2026, 8, 8, 12, 0, tzinfo=UTC)

# Türkçe alfabenin "zor" harfleri + rapor tipografisi. ₺ BİLİNÇLİ olarak yoktur:
# Vera yüzü onu taşımaz ve Faz 1 raporunda parasal tutar basılmaz (ADR-0008).
TURKISH_PROBE = "ığüşöçİĞÜŞÖÇ"


def _dismissed_match() -> MatchRow:
    return MatchRow(
        subject_name="Ahmet Yılmaz",
        subject_role_label="Yönetici",
        caption="AHMET YILMAZOGLU",
        score="0.71",
        decision="dismissed",
        decision_label="Elendi",
        band_label="İnceleme bandı",
        decided_automatically=False,
        flagged=False,
        topics="role.pep",
        datasets="peps",
        source_url="https://example.org/pep/1",
        review_note="Doğum tarihi ve uyruk tutmuyor; farklı kişi.",
        match_features="name: 0.71",
    )


def _confirmed_match() -> MatchRow:
    return MatchRow(
        subject_name=f"Özen Denetim A.Ş. {TURKISH_PROBE}",
        subject_role_label="Şirketin kendisi",
        caption="OZEN DENETIM LTD",
        score="0.93",
        decision="confirmed",
        decision_label="Onaylandı",
        band_label="Otomatik işaretleme bandı",
        decided_automatically=False,
        flagged=True,
        topics="sanction",
        datasets="masak_dondurma",
        source_url="https://example.org/kayit/1",
        review_note="Vergi numarası birebir örtüşüyor.",
        match_features="name: 0.95, taxNumber: 1.00",
    )


@pytest.fixture
def dossier() -> DossierData:
    """Her bölümü dolu, Türkçe karakter içeren örnek dosya."""
    risk = assess_risk(
        screenings=[
            ScreeningView(
                status="completed",
                completed_at=NOW,
                dataset="all",
                matches=(
                    MatchView(
                        subject_name="Özen Denetim A.Ş.",
                        caption="OZEN DENETIM LTD",
                        score=0.93,
                        decision="confirmed",
                        band="auto_flag",
                        flagged=True,
                        decided_automatically=False,
                        topics=("sanction",),
                        datasets=("masak_dondurma",),
                        source_url="https://example.org/kayit/1",
                    ),
                ),
            )
        ],
        discrepancies=[
            DiscrepancyView(
                attribute_label="Vergi no",
                severity="critical",
                status="open",
                declared_value="1234567890",
                observed_value="9876543210",
                observed_source_ref="faaliyet belgesi s.1",
            )
        ],
        assessed_at=NOW,
    )
    section = ScreeningSection(
        screening_id=uuid.uuid4(),
        status_label="Tamamlandı",
        succeeded=True,
        dataset="all",
        algorithm="logic-v1",
        thresholds="otomatik eleme < 0.65; otomatik işaretleme ≥ 0.85",
        started_at="08.08.2026 12:00",
        completed_at="08.08.2026 12:00",
        duration="2.4 sn",
        error=None,
        dataset_versions={"masak_dondurma": {"version": "20260801"}},
        matches=(_confirmed_match(), _dismissed_match()),
    )
    return DossierData(
        entity_id=uuid.uuid4(),
        title=f"Müşteri Tanıma Dosyası — Özen Denetim A.Ş. {TURKISH_PROBE}",
        subject_name="Özen Denetim A.Ş.",
        subject_kind_label="Tüzel kişi",
        generated_at=NOW,
        generated_by="Scryne Şahin",
        organization_name="Yıldız SMMM Bürosu",
        risk=risk,
        profile=(
            ProfileField("Unvan / ad", "Özen Denetim A.Ş."),
            ProfileField("T.C. kimlik no (maskeli)", "*******1234"),
        ),
        related_parties=(
            RelatedParty(
                name="Ahmet Yılmaz",
                rel_type_label="Yönetici",
                role_title="Yönetim Kurulu Başkanı",
                ownership_percent="%51",
                source_label="Belge",
                source_ref="imza sirküleri s.2",
            ),
        ),
        screenings=(section,),
        discrepancies=(
            DiscrepancyRow(
                attribute_label="Vergi no",
                severity_label="Kritik",
                status_label="Açık",
                declared_value="1234567890",
                observed_value="9876543210",
                observed_source_label="Belge",
                observed_source_ref="faaliyet belgesi s.1",
                resolution_note=None,
                reviewed_at="—",
            ),
        ),
        documents=(
            DocumentRow(
                filename="faaliyet-belgesi.pdf",
                doc_type_label="Faaliyet belgesi",
                uploaded_at="08.08.2026 09:00",
                sha256="ab" * 32,
                source_ref=None,
            ),
        ),
        observations=(
            ObservationRow(
                attribute_label="Adres",
                value="Şişli/İstanbul",
                source_label="Beyan",
                source_ref=None,
                observed_at="08.08.2026 09:00",
            ),
        ),
        dataset_versions={
            "masak_dondurma": {
                "version": "20260801",
                "last_export_at": "2026-08-01T03:00:00+00:00",
            }
        },
        screening_ids=(section.screening_id,),
    )


def _normalize(text: str) -> str:
    """Boşluk dizilerini teke indirir.

    PDF'te dar bir sütun metni satır ortasından kırar (``AHMET\\nYILMAZOGLU``); iddia
    mizanpaja değil **içeriğe** olduğundan karşılaştırma normalize metin üzerinden
    yapılır.
    """
    return " ".join(text.split())


def _pdf_text(payload: bytes) -> str:
    reader = PdfReader(io.BytesIO(payload))
    return _normalize("\n".join(page.extract_text() for page in reader.pages))


def _docx_text(payload: bytes) -> str:
    document = DocxDocument(io.BytesIO(payload))
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return _normalize("\n".join(parts))


def _text(payload: bytes) -> str:
    """Üretilen dosyanın metnini biçimden bağımsız olarak çıkarır."""
    return _pdf_text(payload) if payload.startswith(b"%PDF-") else _docx_text(payload)


def test_pdf_uretilir_ve_gecerli_bir_pdf_dir(dossier: DossierData) -> None:
    payload = render_dossier_pdf(dossier)
    assert payload.startswith(b"%PDF-")
    assert len(PdfReader(io.BytesIO(payload)).pages) >= 1


def test_pdf_turkce_karakterleri_korur(dossier: DossierData) -> None:
    """Gömülü Vera yüzü olmadan bu harfler kutucuk olurdu (ADR-0008)."""
    text = _pdf_text(render_dossier_pdf(dossier))
    assert TURKISH_PROBE in text
    assert "Özen Denetim A.Ş." in text
    assert "Şişli/İstanbul" in text


def test_docx_turkce_karakterleri_korur(dossier: DossierData) -> None:
    text = _docx_text(render_dossier_docx(dossier))
    assert TURKISH_PROBE in text
    assert "Şişli/İstanbul" in text


@pytest.mark.parametrize("renderer", [render_dossier_pdf, render_dossier_docx])
def test_elenen_eslesme_gerekcesiyle_raporda_kalir(dossier: DossierData, renderer) -> None:  # type: ignore[no-untyped-def]
    """MASAK savunması "neye bakıldı ve neden elendi"dir; elenenler basılmalıdır."""
    text = _text(renderer(dossier))
    assert "AHMET YILMAZOGLU" in text
    assert "Doğum tarihi ve uyruk tutmuyor" in text


@pytest.mark.parametrize("renderer", [render_dossier_pdf, render_dossier_docx])
def test_risk_seviyesi_ve_gerekceleri_basilir(dossier: DossierData, renderer) -> None:  # type: ignore[no-untyped-def]
    text = _text(renderer(dossier))
    assert "Yüksek" in text
    assert "Onaylanmış yaptırım eşleşmesi" in text
    # Seviyeyi yükseltmeyen ama karar bekleyen bulgu da görünür.
    assert "Kimlik tanımlayıcısında açık tutarsızlık" in text


@pytest.mark.parametrize("renderer", [render_dossier_pdf, render_dossier_docx])
def test_kimlik_numarasi_maskeli_cikar(dossier: DossierData, renderer) -> None:  # type: ignore[no-untyped-def]
    text = _text(renderer(dossier))
    assert "*******1234" in text


@pytest.mark.parametrize("renderer", [render_dossier_pdf, render_dossier_docx])
def test_veri_kumesi_surumleri_ve_esikler_basilir(dossier: DossierData, renderer) -> None:  # type: ignore[no-untyped-def]
    """Rapor "hangi listeye göre, hangi eşikle" sorusunu kendi başına cevaplamalıdır."""
    text = _text(renderer(dossier))
    assert "masak_dondurma" in text
    assert "20260801" in text
    assert "0.85" in text


def test_html_kullanici_verisini_kacirir(dossier: DossierData) -> None:
    """Unvan alanına düşen HTML, şablonda etiket olarak yorumlanmamalıdır."""
    injected = DossierData(
        **{
            **dossier.__dict__,
            "subject_name": "<script>alert(1)</script>",
            "profile": (ProfileField("Unvan / ad", "<script>alert(1)</script>"),),
        }
    )
    html = render_dossier_html(injected)
    assert "<script>alert(1)</script>" not in html
    assert "&lt;script&gt;" in html


def test_bos_dosya_da_uretilebilir() -> None:
    """Hiç taraması/belgesi olmayan varlık için de rapor üretilebilmelidir."""
    risk = assess_risk(screenings=[], assessed_at=NOW)
    data = DossierData(
        entity_id=uuid.uuid4(),
        title="Müşteri Tanıma Dosyası — Yeni Kayıt",
        subject_name="Yeni Kayıt",
        subject_kind_label="Tüzel kişi",
        generated_at=NOW,
        generated_by=None,
        organization_name=None,
        risk=risk,
    )
    text = _pdf_text(render_dossier_pdf(data))
    assert "Belirsiz" in text
    assert "tarama kaydı yok" in text
    assert render_dossier_docx(data).startswith(b"PK")


def test_dataset_version_rows_bozuk_zaman_damgasini_ham_birakir() -> None:
    rows = dataset_version_rows({"x": {"version": "1", "last_export_at": "gecersiz"}})
    assert rows == (("x", "1", "gecersiz"),)
